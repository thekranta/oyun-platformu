# -*- coding: utf-8 -*-
"""Oyunlar ve Maarif Modeli 2026 karşılıkları: pazarlama için Word belgesi (python-docx).

Doğrudan çalıştırmak yerine `node scripts/export-maarif-doc.mjs` kullanın: oyun listesini katalogdan üretir
(scripts/maarif-doc-data.ts) ve bu betiği çağırır. Elle: python scripts/maarif_doc.py <guide.json> <cikti.docx>
Gereksinim: pip install python-docx
"""
import json
import sys
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

if len(sys.argv) != 3:
    print('Kullanım: python scripts/maarif_doc.py <guide.json> <cikti.docx>')
    sys.exit(2)
DATA, OUT = sys.argv[1], sys.argv[2]
data = json.load(open(DATA, encoding='utf-8'))

NAVY = RGBColor(0x1E, 0x4E, 0x8C)
GRAY = RGBColor(0x55, 0x5B, 0x66)
FONT = 'Calibri'

ALAN_COLORS = {
    'Matematik': '3B82F6', 'Fen': '10B981', 'Türkçe': 'F59E0B', 'Sosyal': '14B8A6',
    'Sosyal-Duygusal': 'EC4899', 'Sanat': '8B5CF6', 'Müzik': 'F97316', 'Hareket ve Sağlık': 'EF4444',
}
TIER_STYLE = {  # (zemin, yazı)
    'free': ('E8F5E9', '2E7D32'), 'tohum': ('FFF3E0', 'BF5B00'), 'filiz': ('E3F2FD', '1565C0'),
    'fidan': ('F3E5F5', '6A1B9A'), 'orman': ('F3E5F5', '6A1B9A'),
}
TIER_LABELS = data['tierLabels']


def set_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(attr), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcpr.append(shd)


def cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblpr = table._element.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement('w:' + side)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblpr.append(mar)


def borders(table, color='D5DAE1'):
    tblpr = table._element.tblPr
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + side)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        b.append(el)
    tblpr.append(b)


def fixed_layout(table, widths_cm):
    table.autofit = False
    tblpr = table._element.tblPr
    lay = OxmlElement('w:tblLayout')
    lay.set(qn('w:type'), 'fixed')
    tblpr.append(lay)
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    grid = table._element.find(qn('w:tblGrid'))
    if grid is not None:
        for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
            gc.set(qn('w:w'), str(int(widths_cm[i] / 2.54 * 1440)))


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trpr.append(el)


def cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:cantSplit')
    el.set(qn('w:val'), 'true')
    trpr.append(el)


def write(cell, text, size=9, bold=False, italic=False, color=None, align=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def heading_bar(doc, text, hex_color, sub=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '36')
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), hex_color)
    bdr.append(left)
    ppr.append(bdr)
    r = p.add_run(text)
    set_font(r, size=18, bold=True, color=RGBColor.from_string(hex_color))
    if sub:
        r2 = p.add_run('   ' + sub)
        set_font(r2, size=10, color=GRAY)
    return p


def add_page_number_footer(section, text):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text + '   |   Sayfa ')
    set_font(r, size=8, color=GRAY)
    r2 = p.add_run()
    set_font(r2, size=8, color=GRAY)
    for tag, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if tag:
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), tag)
            r2._r.append(fc)
        else:
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = txt
            r2._r.append(it)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.6)
sec.top_margin = Cm(1.5)
sec.bottom_margin = Cm(1.5)
add_page_number_footer(sec, 'Oyunlar ve Maarif Modeli 2026 karşılıkları')

st = doc.styles['Normal']
st.font.name = FONT
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

# ---------------- Kapak / özet ----------------
months = ['Ocak', 'Şubat', 'Mart', 'Nisan', 'Mayıs', 'Haziran', 'Temmuz', 'Ağustos', 'Eylül', 'Ekim', 'Kasım', 'Aralık']
y, m, d = data['generated'].split('-')
date_text = f"{int(d)} {months[int(m) - 1]} {y}"

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
set_font(p.add_run('Oyunlar ve Maarif Modeli 2026 Karşılıkları'), size=28, bold=True, color=NAVY)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
set_font(p.add_run('Okul Öncesi Eğitim Programı öğrenme çıktıları ile oyun eşleşmesi  ·  ' + date_text), size=12, color=GRAY)

# Sayı kutuları
stats = [
    (str(data['totalGames']), 'oyun'),
    (str(data['totalCodes']), 'kazanım (öğrenme çıktısı)'),
    (str(data['officialCodes']), 'resmi Maarif kodu'),
    (str(len(data['summary'])), 'gelişim alanı'),
]
t = doc.add_table(rows=1, cols=4)
t.alignment = WD_TABLE_ALIGNMENT.LEFT
borders(t, 'E3E8F0')
cell_margins(t, 70, 70, 140, 140)
for i, (num, lab) in enumerate(stats):
    c = t.rows[0].cells[i]
    shade(c, 'F2F6FC')
    pp = c.paragraphs[0]
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(pp.add_run(num), size=26, bold=True, color=NAVY)
    p2 = c.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p2.add_run(lab), size=10, color=GRAY)
fixed_layout(t, [6.6, 6.6, 6.6, 6.6])

doc.add_paragraph().paragraph_format.space_after = Pt(2)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
set_font(p.add_run('Bu belge nasıl okunur? '), size=10.5, bold=True, color=NAVY)
set_font(p.add_run(
    'Her oyun, desteklediği "Türkiye Yüzyılı Maarif Modeli Okul Öncesi Eğitim Programı (2026)" öğrenme çıktısıyla birlikte listelenmiştir. '
    'Çıktı kodu (örn. MAB.4) ve çıktının resmi metni program belgesinden alınmıştır; "Oyun bu kazanımı nasıl çalıştırır?" sütunu '
    'oyunun çocuğa fiilen ne yaptırdığını anlatır. Bölümler gelişim alanına göre dizilmiştir.'), size=10.5)

# Alan özeti + paket etiketleri yan yana (birinci sayfaya sığsın)
outer = doc.add_table(rows=1, cols=2)
outer.alignment = WD_TABLE_ALIGNMENT.LEFT
cell_margins(outer, 0, 0, 0, 160)
left, right = outer.rows[0].cells

hp = left.paragraphs[0]
hp.paragraph_format.space_before = Pt(6)
hp.paragraph_format.space_after = Pt(4)
set_font(hp.add_run('Alanlara göre özet'), size=12, bold=True, color=NAVY)
t = left.add_table(rows=1, cols=3)
borders(t)
cell_margins(t, 35, 35, 90, 90)
hdr = t.rows[0]
for i, h in enumerate(['Gelişim alanı', 'Kazanım', 'Oyun']):
    shade(hdr.cells[i], '1E4E8C')
    write(hdr.cells[i], h, size=9.5, bold=True, color=RGBColor(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER if i else None)
for s_ in data['summary']:
    r = t.add_row()
    cant_split(r)
    col = ALAN_COLORS.get(s_['alan'], '64748B')
    write(r.cells[0], '■  ', size=10, bold=True, color=RGBColor.from_string(col))
    r.cells[0].paragraphs[0].add_run(s_['alan'])
    set_font(r.cells[0].paragraphs[0].runs[-1], size=10, bold=True)
    write(r.cells[1], str(s_['kod']), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    write(r.cells[2], str(s_['oyun']), size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
fixed_layout(t, [6.4, 2.8, 2.8])

hp = right.paragraphs[0]
hp.paragraph_format.space_before = Pt(6)
hp.paragraph_format.space_after = Pt(4)
set_font(hp.add_run('Paket etiketleri'), size=12, bold=True, color=NAVY)
tc = data['tierCount']
legend = [
    ('free', f"{tc.get('free', 0)} oyun", "Ücretsiz pakette açık başlangıç oyunları (Minik Müze'de yalnız vitrin tablosu)."),
    ('tohum', f"{tc.get('tohum', 0)} oyun", 'Tohum ve üzeri paketlerde açık sıradan oyunlar.'),
    ('filiz', f"{tc.get('filiz', 0)} oyun", 'Filiz ve üzeri: Akıllı (performansa göre uyarlanan) oyunlar ve etkileşimli değer hikâyeleri.'),
    ('fidan', f"{tc.get('fidan', 0)} oyun", "Fidan ve üzeri: Müzik Kutusu (40'tan fazla özgün şarkı)."),
]
t2 = right.add_table(rows=0, cols=3)
borders(t2)
cell_margins(t2, 45, 45, 90, 90)
for tier, cnt, desc in legend:
    r = t2.add_row()
    cant_split(r)
    bg, fg = TIER_STYLE[tier]
    shade(r.cells[0], bg)
    write(r.cells[0], TIER_LABELS[tier], size=10, bold=True, color=RGBColor.from_string(fg), align=WD_ALIGN_PARAGRAPH.CENTER)
    write(r.cells[1], cnt, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    write(r.cells[2], desc, size=9.5)
fixed_layout(t2, [2.2, 1.8, 8.6])
fixed_layout(outer, [12.7, 13.4])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
set_font(p.add_run('Not: '), size=9, bold=True, color=GRAY)
set_font(p.add_run(
    'Oyun–kazanım eşleştirmeleri platformun pedagojik değerlendirmesidir ve Milli Eğitim Bakanlığı onayı anlamına gelmez. '
    'Resmi kod bulunmayan sosyal-duygusal becerilerde kod yerine beceri adı gösterilmiştir. '
    'Liste oyun kataloğundan otomatik üretilir; yeni oyunlar eklendikçe güncellenir.'), size=9, color=GRAY)

# ---------------- Alan tabloları ----------------
COLS = [('Oyun', 4.3), ('Maarif kodu', 2.1), ('Resmi kazanım metni', 7.9), ('Oyun bu kazanımı nasıl çalıştırır?', 9.5), ('Paket', 2.4)]
widths = [w for _, w in COLS]
by_alan = {}
for c in data['codes']:
    by_alan.setdefault(c['alan'], []).append(c)
ordered_alans = [s['alan'] for s in data['summary']]

for alan in ordered_alans:
    codes = by_alan[alan]
    n_games = sum(len(c['games']) for c in codes)
    sec_break = doc.add_paragraph()
    sec_break.add_run().add_break(WD_BREAK.PAGE)
    sec_break.paragraph_format.space_after = Pt(0)
    heading_bar(doc, alan, ALAN_COLORS.get(alan, '64748B'), f"{len(codes)} kazanım  ·  {n_games} oyun")

    t = doc.add_table(rows=1, cols=len(COLS))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    borders(t)
    cell_margins(t)
    hdr = t.rows[0]
    repeat_header(hdr)
    cant_split(hdr)
    for i, (h, _) in enumerate(COLS):
        shade(hdr.cells[i], '1E4E8C')
        write(hdr.cells[i], h, size=9.5, bold=True, color=RGBColor(255, 255, 255))

    footnote_needed = False
    for ci, c in enumerate(codes):
        band = 'F7F9FC' if ci % 2 == 0 else 'FFFFFF'
        rows = []
        for g in c['games']:
            r = t.add_row()
            cant_split(r)
            for cell in r.cells:
                shade(cell, band)
            write(r.cells[0], g['name'], size=9.5, bold=True)
            write(r.cells[3], g['detail'] if g['detail'] else 'Kazanımın kendisini doğrudan çalıştırır.', size=9, color=GRAY if not g['detail'] else None)
            bg, fg = TIER_STYLE[g['tier']]
            shade(r.cells[4], bg)
            write(r.cells[4], TIER_LABELS[g['tier']], size=9, bold=True, color=RGBColor.from_string(fg), align=WD_ALIGN_PARAGRAPH.CENTER)
            rows.append(r)
        code_cell = rows[0].cells[1].merge(rows[-1].cells[1]) if len(rows) > 1 else rows[0].cells[1]
        text_cell = rows[0].cells[2].merge(rows[-1].cells[2]) if len(rows) > 1 else rows[0].cells[2]
        for mc in (code_cell, text_cell):
            for extra in mc.paragraphs[1:]:
                extra._element.getparent().remove(extra._element)
            mc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if c['code']:
            write(code_cell, c['code'], size=10, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            write(code_cell, 'Resmi kod yok*', size=8.5, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)
            footnote_needed = True
        write(text_cell, c['text'], size=9.5)
    fixed_layout(t, widths)

    if footnote_needed:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        set_font(p.add_run('* '), size=8.5, bold=True, color=GRAY)
        set_font(p.add_run(
            'Maarif Modeli 2026 Okul Öncesi Programı sosyal-duygusal öğrenme becerilerini isimlendirir ancak kod numarası vermez; '
            'bu nedenle resmi kod yerine beceri adı (köşeli parantez içinde) gösterilmiştir.'), size=8.5, color=GRAY)

doc.core_properties.title = 'Oyunlar ve Maarif Modeli 2026 Karşılıkları'
doc.core_properties.author = 'Oyun Platformu'
doc.save(OUT)
print('kaydedildi', OUT)
